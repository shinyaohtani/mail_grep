"""Tests for hit_report.py - HitReport"""
import pytest
import csv
import sys
from pathlib import Path
from datetime import datetime
import tempfile

sys.path.insert(0, str(Path(__file__).parent.parent))

from hit_report import HitReport
from hit_line import HitLine
from mail_profile import MailProfile


def make_profile(date_dt=None, date_str="", message_id="<test@example.com>"):
    """テスト用MailProfile生成"""
    return MailProfile(
        message_id=message_id,
        date_str=date_str,
        date_dt=date_dt,
        link="message:%3Ctest@example.com%3E",
        subj="Test Subject",
        from_addr="sender@example.com",
        to_addr="receiver@example.com",
    )


def make_hit(profile, mail_id=1, hit_count=1, parttype="text/plain", line="Test"):
    """テスト用HitLine生成"""
    return HitLine(
        mail_keys=profile,
        mail_id=mail_id,
        hit_count=hit_count,
        parttype=parttype,
        line=line,
    )


class TestHitReportBasic:
    """HitReport 基本機能のテスト (4個)"""

    def test_create_and_append(self):
        """レポート作成とヒット行追加"""
        report = HitReport()
        assert len(report.hit_lines) == 0
        profile = make_profile(datetime(2024, 1, 1), "2024-01-01")
        report.append_hit_line(make_hit(profile))
        assert len(report.hit_lines) == 1

    def test_mail_count(self):
        """mail_count - hit_count==1のみカウント"""
        report = HitReport()
        p1 = make_profile(datetime(2024, 1, 1), "2024-01-01", "<mail1@example.com>")
        p2 = make_profile(datetime(2024, 1, 2), "2024-01-02", "<mail2@example.com>")
        report.append_hit_line(make_hit(p1, mail_id=1, hit_count=1))
        report.append_hit_line(make_hit(p1, mail_id=1, hit_count=2))
        report.append_hit_line(make_hit(p2, mail_id=2, hit_count=1))
        assert report.mail_count() == 2

    def test_headers_constant(self):
        """HEADERS定数"""
        assert len(HitReport.HEADERS) == 10
        assert "mail_id" in HitReport.HEADERS
        assert "Matched Line" in HitReport.HEADERS

    def test_mail_count_empty(self):
        """mail_count - 空レポート"""
        report = HitReport()
        assert report.mail_count() == 0


class TestHitReportSort:
    """HitReport.sort() のテスト (2個)"""

    def test_sort_by_date_descending(self):
        """日付降順ソート"""
        report = HitReport()
        p1 = make_profile(datetime(2024, 1, 1), "2024-01-01")
        p2 = make_profile(datetime(2024, 1, 3), "2024-01-03")
        p3 = make_profile(datetime(2024, 1, 2), "2024-01-02")
        report.append_hit_line(make_hit(p1, mail_id=1))
        report.append_hit_line(make_hit(p2, mail_id=2))
        report.append_hit_line(make_hit(p3, mail_id=3))
        report.sort()
        assert report.hit_lines[0].mail_keys.date_str == "2024-01-03"
        assert report.hit_lines[2].mail_keys.date_str == "2024-01-01"

    def test_sort_none_date_last(self):
        """日付なしは末尾"""
        report = HitReport()
        p1 = make_profile(datetime(2024, 1, 1), "2024-01-01")
        p2 = make_profile(None, "")
        report.append_hit_line(make_hit(p1, mail_id=1))
        report.append_hit_line(make_hit(p2, mail_id=2))
        report.sort()
        assert report.hit_lines[-1].mail_keys.date_dt is None


class TestHitReportStoreCsv:
    """HitReport._store_csv() のテスト (3個)"""

    def test_store_csv_creates_file_with_bom(self):
        """CSVファイル作成（BOM付きUTF-8）"""
        report = HitReport()
        profile = make_profile(datetime(2024, 1, 1), "2024-01-01 10:00:00")
        report.append_hit_line(make_hit(profile))

        with tempfile.TemporaryDirectory() as tmpdir:
            csv_path = Path(tmpdir) / "test.csv"
            report.store(csv_path)
            assert csv_path.exists()
            with open(csv_path, "rb") as f:
                assert f.read(3) == b"\xef\xbb\xbf"

    def test_store_csv_header_and_data(self):
        """ヘッダ行とデータ行の確認"""
        report = HitReport()
        profile = make_profile(datetime(2024, 1, 1), "2024-01-01 10:00:00")
        report.append_hit_line(make_hit(profile, line="Test content"))

        with tempfile.TemporaryDirectory() as tmpdir:
            csv_path = Path(tmpdir) / "test.csv"
            report.store(csv_path)
            with open(csv_path, encoding="utf-8-sig") as f:
                reader = csv.reader(f)
                header = next(reader)
                assert header == HitReport.HEADERS
                row = next(reader)
                assert "Test content" in row[-1]

    def test_store_csv_empty_report(self):
        """空レポートの保存"""
        report = HitReport()

        with tempfile.TemporaryDirectory() as tmpdir:
            csv_path = Path(tmpdir) / "test.csv"
            report.store(csv_path)
            with open(csv_path, encoding="utf-8-sig") as f:
                reader = csv.reader(f)
                next(reader)  # header
                assert len(list(reader)) == 0


class TestHitReportStoreXlsx:
    """HitReport._store_xlsx() のテスト (3個)"""

    def test_store_xlsx_creates_file(self):
        """XLSXファイル作成"""
        report = HitReport()
        profile = make_profile(datetime(2024, 1, 1), "2024-01-01 10:00:00")
        report.append_hit_line(make_hit(profile))

        with tempfile.TemporaryDirectory() as tmpdir:
            xlsx_path = Path(tmpdir) / "test.xlsx"
            report.store(xlsx_path)
            assert xlsx_path.exists()

    def test_store_xlsx_valid_format(self):
        """XLSXフォーマットと内容確認"""
        from openpyxl import load_workbook

        report = HitReport()
        profile = make_profile(datetime(2024, 1, 1), "2024-01-01 10:00:00")
        report.append_hit_line(make_hit(profile))

        with tempfile.TemporaryDirectory() as tmpdir:
            xlsx_path = Path(tmpdir) / "test.xlsx"
            report.store(xlsx_path)
            wb = load_workbook(xlsx_path)
            ws = wb.active
            assert ws.title == "results"
            header = [cell.value for cell in ws[1]]
            assert header == HitReport.HEADERS

    def test_store_xlsx_autofilter(self):
        """XLSXのオートフィルタ設定"""
        from openpyxl import load_workbook

        report = HitReport()
        profile = make_profile(datetime(2024, 1, 1), "2024-01-01 10:00:00")
        report.append_hit_line(make_hit(profile))

        with tempfile.TemporaryDirectory() as tmpdir:
            xlsx_path = Path(tmpdir) / "test.xlsx"
            report.store(xlsx_path)
            wb = load_workbook(xlsx_path)
            ws = wb.active
            assert ws.auto_filter.ref is not None
